#!/usr/bin/env python3
"""
xlsx_to_pdf.py — Export one sheet of an Excel workbook to a styled PDF table.

Usage:
    python3 xlsx_to_pdf.py <workbook.xlsx> <output.pdf> [<sheet-name>]

Exit codes:
    0  success
    1  runtime or IO error
    2  bad arguments
    3  required package not installed (openpyxl or reportlab)
"""
import sys
import os

# ── dependency check ──────────────────────────────────────────────────────────
# Exit 3 so the Node.js caller can surface an actionable error message.

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl is required: python3 -m pip install openpyxl", file=sys.stderr)
    sys.exit(3)

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    from reportlab.lib.colors import HexColor
except ImportError:
    print("reportlab is required: python3 -m pip install reportlab", file=sys.stderr)
    sys.exit(3)

# ── palette ───────────────────────────────────────────────────────────────────

HEADER_BG    = HexColor('#2C3E50')
ALT_ROW_BG   = HexColor('#F2F4F5')
GRID_COLOR   = HexColor('#BDC3C7')
BORDER_COLOR = HexColor('#7F8C8D')

# ── helpers ───────────────────────────────────────────────────────────────────

def cell_str(cell):
    """Format a cell value as a display string."""
    v = cell.value
    if v is None:
        return ''
    if isinstance(v, float):
        # Drop the decimal part when the value is a whole number.
        return str(int(v)) if v == int(v) else f'{v:.6g}'
    return str(v)


def read_sheet(ws):
    """
    Read the worksheet into a flat grid suitable for reportlab.

    Returns:
        grid      : list[list[str]]     — every cell as a string; merged
                    non-master cells are set to '' so SPAN commands work.
        spans     : list[((r0,c0),(r1,c1))] — 0-indexed merge ranges
        col_hints : list[float|None]    — openpyxl column widths in character
                    units (None when the sheet has no explicit width set)
    """
    if ws.max_row is None or ws.max_column is None:
        return [], [], []

    nrows, ncols = ws.max_row, ws.max_column

    grid = [['' for _ in range(ncols)] for _ in range(nrows)]

    for row_cells in ws.iter_rows(max_row=nrows, max_col=ncols):
        for cell in row_cells:
            r, c = cell.row - 1, cell.column - 1
            grid[r][c] = cell_str(cell)

    # Collect merge ranges and blank out every non-master cell so reportlab's
    # SPAN command renders correctly.
    spans = []
    for m in ws.merged_cells.ranges:
        r0, c0 = m.min_row - 1, m.min_col - 1
        r1, c1 = m.max_row - 1, m.max_col - 1
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                if (r, c) != (r0, c0) and r < nrows and c < ncols:
                    grid[r][c] = ''
        if r1 > r0 or c1 > c0:
            spans.append(((r0, c0), (r1, c1)))

    col_hints = []
    for i in range(1, ncols + 1):
        dim = ws.column_dimensions.get(get_column_letter(i))
        col_hints.append(dim.width if dim and dim.width else None)

    return grid, spans, col_hints


def build_pdf(grid, spans, col_hints, output_path):
    """Render `grid` as a styled PDF table at `output_path`."""
    if not grid:
        return

    ncols  = max(len(r) for r in grid)
    margin = cm

    # Switch to landscape for wide sheets.
    page_size = landscape(letter) if ncols > 7 else letter
    avail_w   = page_size[0] - 2 * margin

    # Proportional column widths from the openpyxl hints; fall back to equal.
    raw_total = sum(w for w in col_hints[:ncols] if w)
    if raw_total > 0:
        col_widths = [
            avail_w * ((col_hints[i] if i < len(col_hints) and col_hints[i]
                        else raw_total / ncols) / raw_total)
            for i in range(ncols)
        ]
    else:
        col_widths = [avail_w / ncols] * ncols

    doc = SimpleDocTemplate(
        output_path,
        pagesize=page_size,
        leftMargin=margin, rightMargin=margin,
        topMargin=margin,  bottomMargin=margin,
    )

    table = Table(grid, colWidths=col_widths, repeatRows=1)

    style_cmds = [
        # ── header row ────────────────────────────────────────────────────────
        ('BACKGROUND',     (0, 0), (-1, 0),  HEADER_BG),
        ('TEXTCOLOR',      (0, 0), (-1, 0),  colors.white),
        ('FONTNAME',       (0, 0), (-1, 0),  'Helvetica-Bold'),
        ('FONTSIZE',       (0, 0), (-1, 0),  9),
        ('ALIGN',          (0, 0), (-1, 0),  'CENTER'),
        ('VALIGN',         (0, 0), (-1, 0),  'MIDDLE'),
        ('TOPPADDING',     (0, 0), (-1, 0),  6),
        ('BOTTOMPADDING',  (0, 0), (-1, 0),  6),
        ('LINEBELOW',      (0, 0), (-1, 0),  1.5, BORDER_COLOR),
        # ── body rows ─────────────────────────────────────────────────────────
        ('FONTNAME',       (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE',       (0, 1), (-1, -1), 8),
        ('TOPPADDING',     (0, 1), (-1, -1), 3),
        ('BOTTOMPADDING',  (0, 1), (-1, -1), 3),
        ('VALIGN',         (0, 1), (-1, -1), 'TOP'),
        # ── shared padding ────────────────────────────────────────────────────
        ('LEFTPADDING',    (0, 0), (-1, -1), 4),
        ('RIGHTPADDING',   (0, 0), (-1, -1), 4),
        # ── alternating rows ──────────────────────────────────────────────────
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, ALT_ROW_BG]),
        # ── grid and outer border ─────────────────────────────────────────────
        ('GRID',           (0, 0), (-1, -1), 0.5, GRID_COLOR),
        ('BOX',            (0, 0), (-1, -1), 1.0, BORDER_COLOR),
    ]

    # Merged-cell spans: reportlab needs the non-master cells to be '' and a
    # SPAN command per merge range (column-first, row-second indexing).
    for (r0, c0), (r1, c1) in spans:
        style_cmds.append(('SPAN', (c0, r0), (c1, r1)))

    table.setStyle(TableStyle(style_cmds))
    doc.build([table])

# ── entry point ───────────────────────────────────────────────────────────────

def main():
    if len(sys.argv) < 3:
        print(
            f'Usage: {sys.argv[0]} <workbook.xlsx> <output.pdf> [<sheet-name>]',
            file=sys.stderr,
        )
        sys.exit(2)

    xlsx_path   = sys.argv[1]
    output_path = sys.argv[2]
    sheet_name  = sys.argv[3] if len(sys.argv) > 3 else None

    if not os.path.isfile(xlsx_path):
        print(f'File not found: {xlsx_path}', file=sys.stderr)
        sys.exit(1)

    try:
        wb = load_workbook(xlsx_path, data_only=True)
    except Exception as exc:
        print(f'Cannot open workbook: {exc}', file=sys.stderr)
        sys.exit(1)

    if sheet_name:
        if sheet_name not in wb.sheetnames:
            print(
                f'Sheet not found: {sheet_name!r}. '
                f'Available: {wb.sheetnames}',
                file=sys.stderr,
            )
            sys.exit(1)
        ws = wb[sheet_name]
    else:
        ws = wb.active

    try:
        grid, spans, col_hints = read_sheet(ws)
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        build_pdf(grid, spans, col_hints, output_path)
    except Exception as exc:
        print(f'PDF generation failed: {exc}', file=sys.stderr)
        sys.exit(1)


if __name__ == '__main__':
    main()
